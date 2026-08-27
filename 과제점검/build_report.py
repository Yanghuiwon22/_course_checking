"""
build_report.py
───────────────
과제별 최종 통합 리포트 생성기 (Excel 3시트 + TXT).

data/hw{XX}.py 에 두 변수를 채운 뒤 실행:
    python build_report.py --hw 9
    python build_report.py --hw 9 --out ../과제09_점검리포트.xlsx
"""
from __future__ import annotations

import argparse
import importlib.util
import sys
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import BASE_DIR, PROJECT_DIR
from utils import COLORS as C

# ── 색상 매핑 ────────────────────────────────────────────────────
LMS_BG = {"✅": C["green"], "❌": C["pink"], "—": C["gray"]}
GH_BG  = {
    "🟢 이상 없음":   C["green"],
    "🟡 폴더 미정리": C["yellow"],
    "🟡 경고":        C["yellow"],
    "🟠 실행 오류":   C["orange"],
    "🔴 이슈 복수":   C["pink"],
    "⛔ 미제출":      C["gray"],
}
SUM_BG = {"제출": C["green"], "일부제출": C["yellow"], "미제출": C["gray"]}


# ── 데이터 헬퍼 ─────────────────────────────────────────────────
def final_status(lms: str, gh: str) -> str:
    lms_ok = lms in ("✅", "❌")
    gh_ok  = gh != "⛔ 미제출"
    if not lms_ok and not gh_ok: return "미제출"
    if lms == "✅" and gh_ok:     return "제출"
    return "일부제출"


def gh_simple(gh: str) -> str:
    return "—" if gh == "⛔ 미제출" else "✅"


def build_full(students: list) -> list:
    """튜플 인덱스 6에 final_status 삽입 → (이름…로직이슈, 최종종합, *체크값)"""
    return [s[:6] + (final_status(s[3], s[4]),) + s[6:] for s in students]


# ── Excel 셀 헬퍼 ────────────────────────────────────────────────
_thin = Side(border_style="thin", color="CCCCCC")

def _fill(hex_: str)  -> PatternFill: return PatternFill("solid", fgColor=hex_)
def _border()         -> Border:      return Border(**{k: _thin for k in ("left","right","top","bottom")})
def _font(fg="000000", bold=False, size=10) -> Font: return Font(color=fg, bold=bold, size=size)
def _center()         -> Alignment:   return Alignment(horizontal="center", vertical="center", wrap_text=True)
def _lwrap()          -> Alignment:   return Alignment(horizontal="left",   vertical="center", wrap_text=True)

def wcell(ws, r, c, v, *, bg=None, fg="000000", bold=False, sz=10, align=None):
    cl = ws.cell(r, c, v)
    if bg: cl.fill = _fill(bg)
    cl.font      = _font(fg, bold, sz)
    cl.alignment = align or _center()
    cl.border    = _border()
    return cl


# ── 헤더 행 작성 (4행) ──────────────────────────────────────────
BASE_HEADERS = ["이름", "학번", "GitHub ID", "LMS 제출", "GitHub 제출", "로직 이슈", "최종종합"]

def write_headers(ws, check_cols: list[str], last_col: str, hw2: str):
    today = date.today().strftime("%Y-%m-%d")

    # 행1: 제목
    ws.merge_cells(f"A1:{last_col}1")
    wcell(ws, 1, 1, f"과제{hw2} 점검 리포트  ·  {today}", bg=C["dark"], fg=C["white"], bold=True, sz=13)
    ws.row_dimensions[1].height = 30

    # 행2: 부제목
    ws.merge_cells(f"A2:{last_col}2")
    wcell(ws, 2, 1, f"LMS + GitHub 통합 점검", bg=C["navy"], fg=C["white"], bold=True, sz=11)
    ws.row_dimensions[2].height = 24

    # 행3: 그룹 헤더 (A~G 빈 네이비, H~ "코드 체크" 라벤더)
    for col in range(1, 8):
        cl = ws.cell(3, col, "")
        cl.fill = _fill(C["navy"]); cl.border = _border()
    if check_cols:
        ws.merge_cells(f"H3:{last_col}3")
        wcell(ws, 3, 8, "코드 체크", bg=C["lavender"], fg=C["navy"], bold=True)
    ws.row_dimensions[3].height = 20

    # 행4: 열 헤더
    for col, h in enumerate(BASE_HEADERS + check_cols, 1):
        wcell(ws, 4, col, h, bg=C["navy"], fg=C["white"], bold=True)
    ws.row_dimensions[4].height = 40


# ── 학생 데이터 행 작성 ──────────────────────────────────────────
def write_student_row(ws, r: int, s: tuple, *, simple_gh=False):
    """s = (이름, 학번, GitHub_ID, LMS, GitHub, 로직이슈, 최종종합, *체크값...)"""
    stripe = C["stripe"] if r % 2 == 0 else C["white"]

    # A~C: 스트라이프
    for col, val in enumerate(s[:3], 1):
        cl = ws.cell(r, col, val)
        cl.fill = _fill(stripe); cl.font = _font(size=10)
        cl.alignment = _center(); cl.border = _border()

    # D: LMS
    wcell(ws, r, 4, s[3], bg=LMS_BG.get(s[3], C["white"]))

    # E: GitHub (simple_gh=True이면 ✅/— 로만 표시)
    if simple_gh:
        gh_val, gh_bg = gh_simple(s[4]), (C["green"] if s[4] != "⛔ 미제출" else C["gray"])
    else:
        gh_val, gh_bg = s[4], GH_BG.get(s[4], C["white"])
    wcell(ws, r, 5, gh_val, bg=gh_bg)

    # F: 로직이슈
    has_issue = s[5] not in ("이상 없음", "—", "")
    cl = ws.cell(r, 6, s[5])
    cl.fill = _fill(C["pink"] if has_issue else C["white"])
    cl.font = _font(C["issue"] if has_issue else "000000", size=9)
    cl.alignment = _lwrap(); cl.border = _border()

    # G: 최종종합
    wcell(ws, r, 7, s[6], bg=SUM_BG.get(s[6], C["white"]), bold=True)

    # H+: 코드체크
    for offset, checked in enumerate(s[7:]):
        wcell(ws, r, 8 + offset, "✅" if checked else "", bg=C["pink"] if checked else C["white"])

    ws.row_dimensions[r].height = 44


def write_summary(ws, students_full: list, start_row: int, last_col: str):
    counts = {k: sum(1 for s in students_full if s[6] == k) for k in ("제출", "일부제출", "미제출")}
    for label, val in [
        ("📊 제출 현황 요약", None),
        ("전체 학생수",       f"{len(students_full)}명"),
        ("제출",              f"{counts['제출']}명"),
        ("일부제출",          f"{counts['일부제출']}명"),
        ("미제출",            f"{counts['미제출']}명"),
    ]:
        ws.merge_cells(f"A{start_row}:F{start_row}")
        cl = ws.cell(start_row, 1, label)
        cl.fill = _fill(C["sum"]); cl.font = _font(C["white"], bold=True)
        cl.alignment = _center()
        if val is not None:
            wcell(ws, start_row, 7, val, bg=C["sum"], fg=C["white"], bold=True)
        start_row += 1


def set_col_widths(ws, num_check: int):
    widths = [14, 13, 22, 11, 16, 54, 12] + [11] * num_check
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w


# ── 시트 생성 ────────────────────────────────────────────────────
def sheet_main(wb, students_full: list, check_cols: list[str], hw2: str):
    ws = wb.active
    ws.title = "종합현황"
    last_col = get_column_letter(7 + len(check_cols))
    write_headers(ws, check_cols, last_col, hw2)
    for i, s in enumerate(students_full, 5):
        write_student_row(ws, i, s)
    write_summary(ws, students_full, len(students_full) + 6, last_col)
    set_col_widths(ws, len(check_cols))
    ws.freeze_panes = "A5"


def sheet_issues(wb, students_full: list, check_cols: list[str], hw2: str):
    ws = wb.create_sheet("로직이슈상세")
    has_issue = [s for s in students_full if s[5] not in ("이상 없음", "—", "")]
    last_col = get_column_letter(7 + len(check_cols))
    write_headers(ws, check_cols, last_col, hw2)
    for i, s in enumerate(has_issue, 5):
        write_student_row(ws, i, s, simple_gh=True)
    set_col_widths(ws, len(check_cols))
    ws.freeze_panes = "A5"


def sheet_missing(wb, students_full: list, check_cols: list[str], hw2: str):
    ws = wb.create_sheet("미제출")
    missing = [s for s in students_full if s[6] in ("미제출", "일부제출")]
    last_col = get_column_letter(7 + len(check_cols))
    write_headers(ws, check_cols, last_col, hw2)
    for i, s in enumerate(missing, 5):
        write_student_row(ws, i, s, simple_gh=True)
    set_col_widths(ws, len(check_cols))
    ws.freeze_panes = "A5"


# ── TXT 생성 ────────────────────────────────────────────────────
def build_txt(students_full: list, check_cols: list[str], hw2: str, out_path: Path):
    counts = {k: sum(1 for s in students_full if s[6] == k) for k in ("제출", "일부제출", "미제출")}
    lines = [
        f"과제{hw2} 점검 기준  ({date.today():%Y-%m-%d})",
        "=" * 50,
        "",
        "■ 열 구성",
        "  A=이름  B=학번  C=GitHub ID  D=LMS제출  E=GitHub제출  F=로직이슈  G=최종종합",
    ]
    if check_cols:
        lines.append("  코드 체크 열: " + " / ".join(
            f"{'HIJKLMNO'[i]}={h.replace(chr(10),' ')}" for i, h in enumerate(check_cols)
        ))
    lines += [
        "",
        "■ 제출 현황",
        f"  전체 {len(students_full)}명  /  제출 {counts['제출']}명  /  "
        f"일부제출 {counts['일부제출']}명  /  미제출 {counts['미제출']}명",
        "",
        "■ 주의 학생 목록",
    ]
    for s in students_full:
        if s[6] == "제출" and s[5] in ("이상 없음", "—") and not any(s[7:]):
            continue
        checks = [check_cols[i].replace("\n", " ") for i, v in enumerate(s[7:]) if v]
        note_parts = []
        if s[6] != "제출":
            note_parts.append(f"LMS={s[3]} GitHub={s[4]}")
        if s[5] not in ("이상 없음", "—", ""):
            note_parts.append(f"이슈: {s[5]}")
        if checks:
            note_parts.append(f"체크: {', '.join(checks)}")
        lines.append(f"  {s[0]} ({s[1]})  " + "  /  ".join(note_parts))

    out_path.write_text("\n".join(lines), encoding="utf-8")


# ── 진입점 ──────────────────────────────────────────────────────
def run(hw_num: int, out_path: Path | None = None) -> Path:
    hw2 = f"{hw_num:02d}"
    data_path = BASE_DIR / "data" / f"hw{hw2}.py"
    if not data_path.exists():
        raise FileNotFoundError(
            f"data/hw{hw2}.py 없음. data/hw_template.py 복사 후 데이터를 채우세요."
        )

    # data/hwXX.py 동적 로드
    spec = importlib.util.spec_from_file_location(f"hw{hw2}", data_path)
    mod  = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)

    check_cols    = mod.CHECK_COLS
    students_full = build_full(mod.STUDENTS)

    if out_path is None:
        out_path = PROJECT_DIR / f"과제{hw2}_점검리포트.xlsx"
    out_path = Path(out_path)

    wb = openpyxl.Workbook()
    sheet_main(wb, students_full, check_cols, hw2)
    sheet_issues(wb, students_full, check_cols, hw2)
    sheet_missing(wb, students_full, check_cols, hw2)
    wb.save(out_path)

    txt_path = out_path.with_suffix(".txt").with_stem(out_path.stem.replace("리포트", "기준"))
    build_txt(students_full, check_cols, hw2, txt_path)

    print(f"✅ xlsx 저장: {out_path}")
    print(f"✅ txt  저장: {txt_path}")
    return out_path


def main():
    ap = argparse.ArgumentParser(description="과제 통합 리포트 생성")
    ap.add_argument("--hw",  type=int, required=True, help="과제 번호 (예: 9)")
    ap.add_argument("--out", type=str, default=None,  help="출력 xlsx 경로")
    args = ap.parse_args()
    run(args.hw, args.out)


if __name__ == "__main__":
    main()
