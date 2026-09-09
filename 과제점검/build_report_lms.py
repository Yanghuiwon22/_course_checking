"""
build_report_lms.py
────────────────────
GitHub가 없는 과목(스마트팜데이터과학 등)의 LMS 전용 통합 리포트.
과제점검/CLAUDE.md 의 xlsx 디자인 시스템(제목행·그룹헤더·색상·요약)을 따르되
GitHub 관련 열을 전부 제거한 축소판이다.

열 구성: 이름 | 학번 | LMS제출 | 이슈 | 최종종합 | [이슈체크 그룹] | [점수 그룹]

사용법:
    python build_report_lms.py --hw 1 --star-name 주조양 --star-sid 202420921
"""
from __future__ import annotations

import argparse
from datetime import date
from pathlib import Path

import gspread
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from check_hw_rubric import judge_student, find_hw_zip, load_submit_times
from check_lms_content import analyze_student_files
from config import (RUBRICS, ROSTER_PATH, LMS_DIR, LMS_ZIP_DIR, PROJECT_DIR,
                    CREDS_PATH, SPREADSHEET_ID, COL_STUDENT_NAME, COL_STUDENT_ID)
from utils import load_roster_rows, resolve_student_dir, COLORS as C

# 과제별 이슈 체크 열 (헤더, 매칭 문구) — 과제 성격에 따라 다르므로 hw 번호별로 관리한다.
CHECK_COLS_BY_HW = {
    1: [   # hw01: 자기소개 (텍스트 분량/형식 기준)
        ("글자수\n부족", "글자수 부족"),
        ("파일명\n규칙위반", "파일명 규칙 위반"),
        ("지각", "지각 ("),
        ("추출실패\n(수동확인)", "수동 확인 필요"),
    ],
    2: [   # hw02: 파이썬 연습 (함수 정의/호출 기준)
        ("함수정의\n없음", "함수 정의(def) 없음"),
        ("함수호출\n없음", "정의한 함수를 호출한 흔적 없음"),
        ("확장자\n오류", "확장자가"),
        ("필수파일\n없음", "필수 파일 없음"),
        ("지각", "지각 ("),
    ],
}

HW_TITLES = {
    1: "과제01: 자기소개",
    2: "과제02: 파이썬 연습",
}

# AI 의심률 (과제01, 전체 40명 원문 직독 후 사람이 직접 판단 — 자동판정 아님).
# 코드 과제와 달리 에세이는 객관적 판정 근거가 약하므로 50% 미만은 신지 않는다.
AI_SUSPECTS_HW01 = [
    ("이강변", "202410103", "70%",
     "경농 AWS 챗봇(RAG/OpenSearch), 골든플래닛 G-CAMP, Farm Sync, Farmit AI '무당이' 등 "
     "실존 스마트팜 플랫폼을 기능 단위로 상세 나열; 진로 경로도 부서명까지 백과사전식으로 분류"),
    ("조혜윤", "202515457", "65%",
     "4번 항목을 ㄱ)ㄴ)ㄷ) 뉴스 인용 형식으로 작성, 예산액(24억원/4억원)까지 정확히 기재 — "
     "1~3번의 개인적 문체와 대비되는 신문기사체"),
    ("김현준", "202310055", "65%",
     "4번 항목 두 사례 모두 실제 언론사 URL을 출처로 표기(라이브뉴스/연합뉴스), "
     "수치까지 정확(암모니아 20%↓, 전력 35.2%↓) — 뉴스 검색·요약 결과를 그대로 인용한 형태"),
    ("문치웅", "202446401", "60%",
     "'대전시 살수차 운영 시범사업'처럼 잘 알려지지 않은 지자체 사업명과 정확한 수치"
     "(조수입 3%↑, 방제비용 10%↓)를 인용 — 학생이 우연히 알기 어려운 구체성"),
    ("이정제", "202220633", "55%",
     "본문에서 '직업 부분에서 AI를 사용했다'고 직접 밝힘 (5번 항목 한정 자진신고)"),
    ("주조양", "202420921", "50%",
     "'AI와 여러 번 대화하며 직업/자격증 정보를 알게 됐다'고 직접 밝힘 — "
     "단 1~4번은 커리큘럼을 구체적으로 인용하는 등 본인 서술로 보여 전체 의심도는 낮게 책정"),
]

# AI 의심률 (과제02, 전체 40명 코드 직독 후 사람이 직접 판단 — 자동판정 아님).
# 대부분 오타·버그·개인 잡담이 섞인 전형적인 초보자 코드였다.
# 딱 한 명만 과제 범위를 크게 벗어난 명확한 AI 생성 흔적이 있었다.
AI_SUSPECTS_HW02 = [
    ("오태헌", "202310081", "85%",
     "지렁이 게임(OOP 클래스+asyncio 비동기 루프)과 ipywidgets 위젯 사용법을 "
     "'## 기능/장점' 구조의 챗봇 특유 설명체로 작성 — '함수 작성/호출 연습' 범위를 "
     "크게 벗어나며, 영어 설명 문장 구조까지 AI 어시스턴트 응답과 동일한 패턴"),
]

# 우수학생 (사람이 원문을 직접 읽고 선정 — AI_SUSPECTS 명단에 오른 학생은 제외).
STAR_STUDENTS_HW01 = [
    ("김선우", "202321613",
     "3491자로 최장문 — 5개 항목 모두 커리큘럼(Pandas/NumPy/KNN 등)과 연결해 구체적으로 서술"),
    ("이성학", "202210099",
     "여름 스마트팜 실습 경험(적엽 시기, 적과 등 실제 현장 용어)을 항목마다 구체적으로 연결"),
    ("조현우", "202310101",
     "편지 형식이 자연스럽고 목표·사례·진로 항목 간 연결이 매끄러움"),
]

STAR_STUDENTS_HW02 = [
    ("문치웅", "202446401",
     "메뉴 선택형 단위변환 프로그램·BMI 계산기·두 점 거리 계산까지 스스로 확장 구현. "
     "'코랩 input()에서 KeyboardInterrupt 발생 원인을 검색해 알아냄' 등 직접 디버깅한 기록도 남김"),
    ("서연우", "202310070",
     "근의 공식 프로그램에 'a=0이면 1차방정식', '판별식<0이면 실근 없음' 예외 처리를 "
     "직접 추가하고, 왜 그렇게 만들었는지 서술까지 남김 — 과제 예시를 그대로 베끼지 않고 보완함"),
]

STAR_STUDENTS_BY_HW = {1: STAR_STUDENTS_HW01, 2: STAR_STUDENTS_HW02}
AI_SUSPECTS_BY_HW    = {1: AI_SUSPECTS_HW01,   2: AI_SUSPECTS_HW02}


# 셀 헬퍼
_thin = Side(border_style="thin", color="CCCCCC")


def _fill(hex_):
    return PatternFill("solid", fgColor=hex_)


def _border():
    sides = {"left": _thin, "right": _thin, "top": _thin, "bottom": _thin}
    return Border(**sides)


def _center():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)


def _lwrap():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)


def wcell(ws, r, c, v, bg=None, fg="000000", bold=False, sz=10, align=None):
    cell = ws.cell(r, c, v)
    cell.font = Font(color=fg, bold=bold, size=sz)
    cell.alignment = align if align else _center()
    cell.border = _border()
    if bg:
        cell.fill = _fill(bg)
    return cell


def collect(hw_num):
    rubric = RUBRICS[hw_num]
    roster = load_roster_rows(ROSTER_PATH, COL_STUDENT_NAME, COL_STUDENT_ID)
    submit_times = load_submit_times(find_hw_zip(LMS_ZIP_DIR, hw_num))
    folder_name = "hw{:02d}".format(hw_num)

    rows = []
    for name, sid in roster:
        student_dir = resolve_student_dir(LMS_DIR, name, sid)
        hw_dir = student_dir / folder_name
        files = analyze_student_files(hw_dir)
        top = None
        for k in submit_times:
            if k.startswith(name):
                top = k
                break
        submit_time = submit_times.get(top)
        verdict, issues = judge_student(name, sid, files, hw_dir, rubric, submit_time)

        if verdict == "미제출":
            lms = "미제출"
        elif submit_time and rubric.get("deadline") and submit_time > rubric["deadline"]:
            lms = "지각"
        else:
            lms = "-"

        rows.append({
            "name": name,
            "sid": sid,
            "lms": lms,
            "issue_text": "; ".join(issues) if issues else "이상없음",
            "verdict": verdict,
            "issues": issues,
        })
    return rows


def score_row(row, star, check_cols):
    if row["verdict"] == "미제출":
        return 0.0, [False] * len(check_cols)

    checks = []
    for _, marker in check_cols:
        hit = False
        for issue in row["issues"]:
            if marker in issue:
                hit = True
                break
        checks.append(hit)

    score = 5.0 - 0.5 * sum(checks)
    if star:
        score += 1.0
    return score, checks


BASE_HEADERS = ["이름", "학번", "LMS\n제출", "이슈", "최종\n종합"]
SCORE_HEADERS = ["기본\n점수", "우수\n학생", "최종"]


def write_headers(ws, hw2, subtitle, last_col, check_cols):
    today = date.today().strftime("%Y-%m-%d")
    n_base = len(BASE_HEADERS)
    n_check = len(check_cols)

    ws.merge_cells("A1:{}1".format(last_col))
    wcell(ws, 1, 1, "\U0001F4CB 프원실(2026) 과제{} 점검 리포트 (LMS)".format(hw2),
          bg=C["dark"], fg=C["white"], bold=True, sz=13)
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:{}2".format(last_col))
    wcell(ws, 2, 1, "{}  |  점검일: {}  |  제출 5 | 오류 -0.5 | 우수 1".format(subtitle, today),
          bg=C["navy"], fg=C["white"], bold=True, sz=11)
    ws.row_dimensions[2].height = 24

    for col in range(1, n_base + 1):
        c = ws.cell(3, col, "")
        c.fill = _fill(C["navy"])
        c.border = _border()

    check_start = n_base + 1
    check_end = n_base + n_check
    ws.merge_cells("{}3:{}3".format(get_column_letter(check_start), get_column_letter(check_end)))
    wcell(ws, 3, check_start, "이슈 체크", bg=C["lavender"], fg=C["navy"], bold=True)

    score_start = check_end + 1
    ws.merge_cells("{}3:{}3".format(get_column_letter(score_start), last_col))
    wcell(ws, 3, score_start, "점수", bg=C["lavender"], fg=C["navy"], bold=True)
    ws.row_dimensions[3].height = 20

    headers = BASE_HEADERS + [h for h, _ in check_cols] + SCORE_HEADERS
    for col, h in enumerate(headers, 1):
        wcell(ws, 4, col, h, bg=C["navy"], fg=C["white"], bold=True)
    ws.row_dimensions[4].height = 40


LMS_BG = {"-": C["green"], "지각": C["yellow"], "미제출": C["gray"]}
VERDICT_BG = {"이상없음": C["green"], "이슈": C["yellow"], "미제출": C["gray"]}


def disp(value):
    """'이상없음'은 프로젝트 관례대로 '-' 로 표시한다 (내부 판정 값은 그대로 유지)."""
    return "-" if value == "이상없음" else value


def write_row(ws, r, row, star, check_cols):
    stripe = C["stripe"] if r % 2 == 0 else C["white"]
    name_label = row["name"] + (" ⭐" if star else "")
    wcell(ws, r, 1, name_label, bg=stripe, sz=10)
    wcell(ws, r, 2, row["sid"], bg=stripe, sz=10)
    wcell(ws, r, 3, row["lms"], bg=LMS_BG.get(row["lms"], C["white"]))

    has_issue = row["issue_text"] != "이상없음"
    cell = ws.cell(r, 4, disp(row["issue_text"]))
    cell.fill = _fill(C["pink"] if has_issue else C["white"])
    cell.font = Font(color=C["issue"] if has_issue else "000000", size=9)
    cell.alignment = _lwrap()
    cell.border = _border()

    final_label = disp(row["verdict"])
    if star and row["verdict"] == "이상없음":
        final_label = "- ⭐"
    wcell(ws, r, 5, final_label, bg=VERDICT_BG.get(row["verdict"], C["white"]), bold=True)

    score, checks = score_row(row, star, check_cols)
    for i, checked in enumerate(checks):
        wcell(ws, r, 6 + i, "✅" if checked else "—",
              bg=C["pink"] if checked else C["white"])

    star_col = 6 + len(check_cols)
    base_score = 0.0 if row["verdict"] == "미제출" else 5.0
    wcell(ws, r, star_col, base_score)
    wcell(ws, r, star_col + 1, 1 if star else "")
    wcell(ws, r, star_col + 2, score, bold=True)
    ws.row_dimensions[r].height = 44


def write_summary(ws, rows, star_count, start_row, last_col):
    counts = {"이상없음": 0, "이슈": 0, "미제출": 0}
    for r in rows:
        counts[r["verdict"]] = counts.get(r["verdict"], 0) + 1

    ws.merge_cells("A{0}:{1}{0}".format(start_row, last_col))
    text = ("\U0001F4CA 제출 현황 요약  |  전체 {}명  |  이상없음 {}명  |  "
           "이슈 {}명  |  미제출 {}명  |  우수학생 {}명").format(
        len(rows), counts["이상없음"], counts["이슈"], counts["미제출"], star_count)
    wcell(ws, start_row, 1, text, bg=C["sum"], fg=C["white"], bold=True, sz=11)
    ws.row_dimensions[start_row].height = 26


def write_ai_suspects(ws, start_row, last_col, suspects):
    """AI 의심률 표를 하단에 추가한다. 사람이 원문을 직접 읽고 판단한 결과이며
    자동 판정이 아니다 — 에세이 특성상 코드보다 근거가 약해 50% 미만은 신지 않는다."""
    r = start_row
    ws.merge_cells("A{0}:{1}{0}".format(r, last_col))
    wcell(ws, r, 1, "⚠ AI 의심률 (사람이 원문을 직접 읽고 판단 — 자동판정 아님, 50% 이상만 기재)",
          bg=C["lavender"], fg=C["navy"], bold=True, sz=10)
    r += 1

    headers = ["학생", "AI 의심률", "사유"]
    for c, h in enumerate(headers, 1):
        wcell(ws, r, c, h, bg=C["navy"], fg=C["white"], bold=True)
    if last_col != "C":
        ws.merge_cells("C{0}:{1}{0}".format(r, last_col))
    r += 1

    for name, sid, pct, reason in suspects:
        wcell(ws, r, 1, "{} ({})".format(name, sid), align=_lwrap(), sz=10)
        wcell(ws, r, 2, pct, bold=True)
        cell = ws.cell(r, 3, reason)
        cell.alignment = _lwrap()
        cell.border = _border()
        cell.font = Font(size=9)
        if last_col != "C":
            ws.merge_cells("C{0}:{1}{0}".format(r, last_col))
        ws.row_dimensions[r].height = 32
        r += 1
    return r


def write_star_reasons(ws, start_row, last_col, star_reasons):
    """우수학생 선발 이유 표 (AI 의심률 표 바로 아래에 놓인다)."""
    r = start_row
    ws.merge_cells("A{0}:{1}{0}".format(r, last_col))
    wcell(ws, r, 1, "⭐ 우수학생 선발 이유 (사람이 원문을 직접 읽고 선정)",
          bg=C["green"], fg=C["navy"], bold=True, sz=10)
    r += 1

    headers = ["학생", "선발 이유"]
    wcell(ws, r, 1, headers[0], bg=C["navy"], fg=C["white"], bold=True)
    wcell(ws, r, 2, headers[1], bg=C["navy"], fg=C["white"], bold=True)
    if last_col not in ("A", "B"):
        ws.merge_cells("B{0}:{1}{0}".format(r, last_col))
    r += 1

    for name, sid, reason in star_reasons:
        wcell(ws, r, 1, "{} ({})".format(name, sid), align=_lwrap(), sz=10)
        cell = ws.cell(r, 2, reason)
        cell.alignment = _lwrap()
        cell.border = _border()
        cell.font = Font(size=9)
        if last_col not in ("A", "B"):
            ws.merge_cells("B{0}:{1}{0}".format(r, last_col))
        ws.row_dimensions[r].height = 32
        r += 1
    return r


def run(hw_num, out_path=None, star_keys=None, ai_suspects=None, star_reasons=None):
    hw2 = "{:02d}".format(hw_num)
    rows = collect(hw_num)
    star_keys = set(star_keys or ())
    check_cols = CHECK_COLS_BY_HW[hw_num]

    last_col = get_column_letter(len(BASE_HEADERS) + len(check_cols) + len(SCORE_HEADERS))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "종합현황"
    subtitle = "2026년 1학기  |  스마트팜데이터과학  |  {}".format(HW_TITLES[hw_num])
    write_headers(ws, hw2, subtitle, last_col, check_cols)

    r = 5
    star_count = 0
    for row in rows:
        is_star = (row["name"], row["sid"]) in star_keys
        if is_star:
            star_count += 1
        write_row(ws, r, row, is_star, check_cols)
        r += 1
    summary_row = r + 1
    write_summary(ws, rows, star_count, summary_row, last_col)

    next_row = summary_row + 2
    if ai_suspects:
        next_row = write_ai_suspects(ws, next_row, last_col, ai_suspects) + 1
    if star_reasons:
        write_star_reasons(ws, next_row, last_col, star_reasons)

    ws.freeze_panes = "A5"
    widths = [14, 12, 10, 54, 12] + [10] * len(check_cols) + [10, 8, 8]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    if out_path is None:
        out_path = PROJECT_DIR / "과제{}_점검리포트.xlsx".format(hw2)
    wb.save(out_path)
    print("xlsx 저장: {}".format(out_path))
    return Path(out_path)


def _hex_to_rgb(h):
    h = h.lstrip("#")
    return {"red": int(h[0:2], 16) / 255, "green": int(h[2:4], 16) / 255, "blue": int(h[4:6], 16) / 255}


def add_score_chart(sh, ws, n_rows, last_col_idx0, anchor_row_1based):
    """'이름'(A열)과 '최종'(마지막 열) 점수를 막대 차트로 추가해 학생별 점수 추이를 보여준다.
    재실행 시 중복 생성되지 않도록 이 시트의 기존 차트를 먼저 지운다."""
    meta = sh.fetch_sheet_metadata()
    del_requests = []
    for s in meta["sheets"]:
        if s["properties"]["sheetId"] != ws.id:
            continue
        for ch in s.get("charts", []):
            del_requests.append({"deleteEmbeddedObject": {"objectId": ch["chartId"]}})
    if del_requests:
        sh.batch_update({"requests": del_requests})

    data_start_0 = 1                  # 헤더(0행) 다음부터
    data_end_0 = 1 + n_rows

    def col_range(col_idx):
        return {"sourceRange": {"sources": [{
            "sheetId": ws.id, "startRowIndex": data_start_0, "endRowIndex": data_end_0,
            "startColumnIndex": col_idx, "endColumnIndex": col_idx + 1}]}}

    chart_request = {"addChart": {"chart": {
        "spec": {
            "title": "학생별 최종 점수 추이",
            "basicChart": {
                "chartType": "COLUMN",
                "legendPosition": "NO_LEGEND",
                "axis": [
                    {"position": "BOTTOM_AXIS", "title": "이름"},
                    {"position": "LEFT_AXIS", "title": "최종 점수"},
                ],
                "domains": [{"domain": col_range(0)}],           # A열: 이름
                "series": [{"series": col_range(last_col_idx0)}], # 마지막 열: 최종
            },
        },
        "position": {"overlayPosition": {
            "anchorCell": {"sheetId": ws.id, "rowIndex": anchor_row_1based - 1, "columnIndex": 0},
            "widthPixels": 900, "heightPixels": 320,
        }},
    }}}
    sh.batch_update({"requests": [chart_request]})


def write_google_sheet(hw_num, hw2, rows, star_keys, ai_suspects=None, star_reasons=None):
    """같은 스프레드시트에 '과제{NN}' 탭을 xlsx와 동일한 열 구성으로 다시 쓴다."""
    check_cols = CHECK_COLS_BY_HW[hw_num]
    sheet_header = (["이름", "학번", "LMS제출", "이슈", "최종종합"]
                    + [h.replace("\n", "") for h, _ in check_cols]
                    + ["기본점수", "우수학생", "최종"])

    gc = gspread.service_account(filename=str(CREDS_PATH))
    sh = gc.open_by_key(SPREADSHEET_ID)
    title = "과제{}".format(hw2)
    star_keys = set(star_keys or ())

    try:
        ws = sh.worksheet(title)
        ws.clear()
    except gspread.WorksheetNotFound:
        ws = sh.add_worksheet(title=title, rows=len(rows) + 30, cols=len(sheet_header))

    # ws.clear()는 값만 지우고 그리드 크기·병합은 그대로 둔다.
    # 예전에 더 적은 열/행으로 만들어진 탭이 남아 있으면, 새로 늘어난 열(예: '최종')이나
    # 행(AI의심률/우수학생 표)에 쓴 값이 그리드 밖으로 밀려 조용히 사라질 수 있으므로
    # 필요한 크기보다 작으면 먼저 키운다.
    needed_cols = len(sheet_header)
    needed_rows = len(rows) + 30   # 본문 + 빈 줄 + AI의심률 표 + 우수학생 표 + 차트 요약
    if ws.col_count < needed_cols or ws.row_count < needed_rows:
        ws.resize(rows=max(ws.row_count, needed_rows), cols=max(ws.col_count, needed_cols))

    # 값만 지운다고 기존 병합까지 없어지지 않으므로, 매번 전체 병합을 해제하고 시작한다.
    sh.batch_update({"requests": [{"unmergeCells": {"range": {"sheetId": ws.id}}}]})

    values = [sheet_header]
    star_count = 0
    row_meta = []   # (verdict, lms) 색칠용
    for row in rows:
        is_star = (row["name"], row["sid"]) in star_keys
        if is_star:
            star_count += 1
        score, checks = score_row(row, is_star, check_cols)
        base_score = 0.0 if row["verdict"] == "미제출" else 5.0
        final_verdict = disp(row["verdict"])
        if is_star and row["verdict"] == "이상없음":
            final_verdict = "- ⭐"
        values.append([
            row["name"] + (" ⭐" if is_star else ""), row["sid"], row["lms"],
            disp(row["issue_text"]), final_verdict,
            *["✅" if c else "—" for c in checks],
            base_score, 1 if is_star else "", score,
        ])
        row_meta.append((row["verdict"], row["lms"]))

    last_col_letter = get_column_letter(len(sheet_header))

    def pad_to(target_row):
        values.extend([[]] * (target_row - len(values) - 1))

    ai_start_row = None
    if ai_suspects:
        ai_start_row = len(values) + 2   # 표 아래 빈 줄 하나 두고 시작
        pad_to(ai_start_row)
        values.append(["⚠ AI 의심률 (사람이 원문을 직접 읽고 판단 — 자동판정 아님, 50% 이상만 기재)"])
        values.append(["학생", "AI 의심률", "사유"])
        for name, sid, pct, reason in ai_suspects:
            values.append(["{} ({})".format(name, sid), pct, reason])

    star_start_row = None
    if star_reasons:
        star_start_row = len(values) + 2
        pad_to(star_start_row)
        values.append(["⭐ 우수학생 선발 이유 (사람이 원문을 직접 읽고 선정)"])
        values.append(["학생", "선발 이유"])
        for name, sid, reason in star_reasons:
            values.append(["{} ({})".format(name, sid), reason])

    ws.update(values)

    fmt_requests = [{
        "range": "A1:{}1".format(last_col_letter),
        "format": {"backgroundColor": {"red": 0.18, "green": 0.25, "blue": 0.34},
                  "textFormat": {"bold": True, "foregroundColor": {"red": 1, "green": 1, "blue": 1}}},
    }]
    for i, (verdict, lms) in enumerate(row_meta, start=2):
        bg = VERDICT_BG.get(verdict)
        if bg:
            fmt_requests.append({"range": "E{}".format(i), "format": {"backgroundColor": _hex_to_rgb(bg)}})
        bg2 = LMS_BG.get(lms)
        if bg2:
            fmt_requests.append({"range": "C{}".format(i), "format": {"backgroundColor": _hex_to_rgb(bg2)}})

    navy = {"red": 0.18, "green": 0.25, "blue": 0.34}
    white = {"red": 1, "green": 1, "blue": 1}
    if ai_start_row:
        fmt_requests.append({
            "range": "A{0}:{1}{0}".format(ai_start_row, last_col_letter),
            "format": {"backgroundColor": {"red": 0.84, "green": 0.79, "blue": 0.94},
                      "textFormat": {"bold": True, "foregroundColor": {"red": 0.18, "green": 0.25, "blue": 0.34}}},
        })
        fmt_requests.append({
            "range": "A{0}:C{0}".format(ai_start_row + 1),
            "format": {"backgroundColor": navy, "textFormat": {"bold": True, "foregroundColor": white}},
        })

    if star_start_row:
        fmt_requests.append({
            "range": "A{0}:{1}{0}".format(star_start_row, last_col_letter),
            "format": {"backgroundColor": _hex_to_rgb(C["green"]),
                      "textFormat": {"bold": True, "foregroundColor": {"red": 0.18, "green": 0.25, "blue": 0.34}}},
        })
        fmt_requests.append({
            "range": "A{0}:B{0}".format(star_start_row + 1),
            "format": {"backgroundColor": navy, "textFormat": {"bold": True, "foregroundColor": white}},
        })

    ws.batch_format(fmt_requests)   # 한 번의 API 호출로 모든 서식 적용 (쓰기 할당량 절약)

    ws.freeze(rows=1)

    chart_anchor_row = len(values) + 3
    add_score_chart(sh, ws, len(rows), len(sheet_header) - 1, chart_anchor_row)

    print("구글시트 업데이트: {} 탭 ({}행, 우수학생 {}명, 차트 추가)".format(title, len(rows), star_count))


def main():
    ap = argparse.ArgumentParser(description="LMS 전용 통합 리포트 생성")
    ap.add_argument("--hw", type=int, required=True)
    ap.add_argument("--out", type=str, default=None)
    ap.add_argument("--star-name", type=str, default=None, help="최고 학생 이름 (직접 지정 시 1명만)")
    ap.add_argument("--star-sid", type=str, default=None, help="최고 학생 학번")
    ap.add_argument("--sheet", action="store_true", help="구글시트 '과제NN' 탭도 같은 형식으로 갱신")
    args = ap.parse_args()
    if args.star_name:
        star_reasons = [(args.star_name, args.star_sid, "직접 지정")]
    else:
        star_reasons = STAR_STUDENTS_BY_HW.get(args.hw, [])
    stars = {(n, s) for n, s, _ in star_reasons}
    ai_suspects = AI_SUSPECTS_BY_HW.get(args.hw)
    run(args.hw, args.out, stars, ai_suspects, star_reasons)
    if args.sheet:
        hw2 = "{:02d}".format(args.hw)
        write_google_sheet(args.hw, hw2, collect(args.hw), stars, ai_suspects, star_reasons)


if __name__ == "__main__":
    main()
