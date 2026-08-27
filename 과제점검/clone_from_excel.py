"""
Clone/pull student GitHub repos based on column M in 프원실(2026)_과제확인.xlsx.

Usage:
    python clone_from_excel.py [--dry-run]
"""

from __future__ import annotations

from pathlib import Path
import argparse
import os
import re
import shutil
import subprocess

import openpyxl

from config import ROSTER_PATH, GITHUB_DIR, COL_STUDENT_NAME, COL_CLONE_CMD
from utils import normalize_name

EXCEL_PATH  = ROSTER_PATH
OUTPUT_ROOT = GITHUB_DIR

# Column indices (1-based)
COL_NAME  = COL_STUDENT_NAME
COL_CLONE = COL_CLONE_CMD


def extract_clone_url(cell_val: str) -> str | None:
    if not cell_val:
        return None
    s = str(cell_val).strip()
    m = re.search(r"git\s+clone\s+(\S+)", s)
    url = m.group(1).strip() if m else s
    if not url or url.endswith("github.com//"):
        return None
    return url


def remove_invalid_entries(target: Path) -> None:
    """'...' 처럼 Windows에서 문제가 되는 이름의 파일/폴더를 제거한다."""
    for root, dirs, files in os.walk(str(target), topdown=False):
        for name in dirs + files:
            if name == "...":
                full = os.path.join(root, name)
                try:
                    if os.path.isdir(full):
                        shutil.rmtree(full, ignore_errors=True)
                    else:
                        os.remove(full)
                    print(f"  [clean] 제거: {full}")
                except Exception as e:
                    print(f"  [clean] 제거 실패: {full} ({e})")


def run(cmd: list[str], dry_run: bool) -> int:
    if dry_run:
        print("[dry-run]", " ".join(cmd))
        return 0
    return subprocess.call(cmd)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true", help="Print actions without executing git commands.")
    args = ap.parse_args()

    if not EXCEL_PATH.exists():
        print(f"엑셀 파일 없음: {EXCEL_PATH}")
        return 1

    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb.active

    total = success = skipped = failed = 0

    for r in range(2, ws.max_row + 1):
        name_raw = ws.cell(row=r, column=COL_NAME).value
        clone_raw = ws.cell(row=r, column=COL_CLONE).value
        if not name_raw and not clone_raw:
            continue

        total += 1
        name = normalize_name(name_raw or f"row{r}")
        url = extract_clone_url(clone_raw)
        if not url:
            print(f"[skip] {name}: clone URL 없음/오류")
            skipped += 1
            continue

        target = OUTPUT_ROOT / name
        if target.exists():
            if (target / ".git").is_dir():
                print(f"[pull] {name}")
                if args.dry_run:
                    print(f"[dry-run] git -C {target} fetch origin")
                    print(f"[dry-run] git -C {target} reset --hard origin/HEAD")
                    code = 0
                else:
                    # 이전 실행이 중단되어 남은 lock 파일 전체 제거
                    git_dir = target / ".git"
                    lock_files = list(git_dir.rglob("*.lock"))
                    for lf in lock_files:
                        try:
                            lf.unlink()
                            print(f"  [lock] {lf.name} 제거: {name}")
                        except Exception as e:
                            print(f"  [lock] {lf.name} 제거 실패: {e}")
                    subprocess.call(["git", "-C", str(target), "config", "core.protectNTFS", "false"])
                    subprocess.call(["git", "-C", str(target), "fetch", "origin"])
                    remove_invalid_entries(target)
                    code = subprocess.call(["git", "-C", str(target), "reset", "--hard", "origin/HEAD"])
                if code == 0:
                    success += 1
                else:
                    print(f"[fail] pull 실패: {name}")
                    failed += 1
            else:
                print(f"[skip] {name}: 폴더 존재하지만 git repo 아님")
                skipped += 1
            continue

        print(f"[clone] {name}")
        code = run(["git", "clone", "-c", "core.protectNTFS=false", url, str(target)], args.dry_run)
        if code == 0:
            success += 1
        else:
            print(f"[fail] clone 실패: {name}")
            failed += 1

    wb.close()
    print(f"완료: 총 {total}, 성공 {success}, 스킵 {skipped}, 실패 {failed}")
    return 0 if failed == 0 else 2


if __name__ == "__main__":
    raise SystemExit(main())
